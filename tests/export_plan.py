"""Run: python tests/export_plan.py — writes tests/test_plan.xlsx for review."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tests.metrics import RUN_SUMMARY_METRICS, STEP_METRICS  # noqa: E402
from tests.scenarios import SAMPLE_DATA, SCENARIOS  # noqa: E402

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
PROFILE_LABEL = {"new": "New user (not onboarded)", "ready": "Onboarded, no history",
                 "history": "Onboarded + 4 weeks sample data"}


def style_sheet(ws, widths, header_row=1):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row == header_row, color="FFFFFF" if c.row == header_row else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.row == header_row:
                c.fill = HEAD_FILL
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def checks_text(step):
    parts = []
    if step.get("intent"):
        parts.append(f"Intent = {step['intent']}")
    if step.get("tools"):
        parts.append("Must call: " + ", ".join(step["tools"]))
    if step.get("no_tools"):
        parts.append("Must NOT call: " + ", ".join(step["no_tools"]))
    for d in step.get("data", []):
        parts.append(f"DB: {d}")
    for r in step.get("reply_has", []):
        parts.append(f"Reply matches: {r}")
    for r in step.get("reply_not", []):
        parts.append(f"Reply must not match: {r}")
    if step.get("judge"):
        parts.append("LLM judge score ≥ 4/5")
    return "\n".join(parts) or "No crash (setup step)"


wb = Workbook()

# Read me
ws = wb.active
ws.title = "Read me"
total_steps = sum(len(s["steps"]) for s in SCENARIOS)
multi = sum(1 for s in SCENARIOS if len(s["steps"]) > 1)
rows = [
    ["Spotter evaluation plan"],
    [""],
    ["Scenarios", len(SCENARIOS)],
    ["Messages (steps)", total_steps],
    ["Multi-turn scenarios", multi],
    ["Categories", ", ".join(dict.fromkeys(s["category"] for s in SCENARIOS))],
    [""],
    ["How a run works", "Each scenario gets its own throwaway test user, starting from the profile in the 'Profile' column. "
                        "Messages go through the same engine as the simulator. After every message the runner records the "
                        "metrics in 'Metrics tracked' and runs the checks in 'Pass criteria'. Test users are deleted afterwards."],
    ["Pass rule", "A step passes only if every check in its 'Pass criteria' passes, plus the style checks that apply to every "
                  "reply (≤1 emoji outside confirmations/reports, no markdown ** or #, at most one question)."],
    ["Results file", "tests/results/eval_<date>.xlsx with sheets: Summary, Steps (one row per message), LLM calls, Errors."],
    ["Model", "gemini-3.5-flash-lite (free tier: 15 requests/min, 500/day). The run is paced to stay under the limit."],
    ["Not built yet", "RPT-03 to RPT-07 need two additions before running: 'weeks_ago' for the weekly report, and a "
                      "get_progress tool for period stats (diet days met, sessions, lift trend, averages)."],
]
for r in rows:
    ws.append(r)
ws["A1"].font = Font(name=FONT, size=14, bold=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=FONT, size=10, bold=c.column == 1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 110

# Metrics tracked
ws = wb.create_sheet("Metrics tracked")
ws.append(["#", "Metric (column in results)", "What it means", "Source"])
for i, (name, desc, src) in enumerate(STEP_METRICS, start=1):
    ws.append([i, name, desc, src])
style_sheet(ws, [5, 26, 80, 30])

# Run summary
ws = wb.create_sheet("Run summary metrics")
ws.append(["Metric", "What it shows"])
for name, desc in RUN_SUMMARY_METRICS:
    ws.append([name, desc])
style_sheet(ws, [36, 80])

# Test cases
ws = wb.create_sheet("Test cases")
ws.append(["Scenario ID", "Category", "Scenario", "Profile", "Multi-turn", "Step", "User message / action",
           "Expected behaviour", "Pass criteria", "Setup"])
for s in SCENARIOS:
    for n, step in enumerate(s["steps"], start=1):
        msg = step.get("say") or f"[system] {step.get('action')}"
        if len(msg) > 160:
            msg = msg[:150] + f"… ({len(msg)} chars)"
        ws.append([s["id"], s["category"], s["name"], PROFILE_LABEL[s["profile"]],
                   "Yes" if len(s["steps"]) > 1 else "No", n, msg, step.get("expect", ""), checks_text(step),
                   s.get("setup", "") if n == 1 else ""])
style_sheet(ws, [11, 16, 30, 22, 9, 6, 46, 46, 46, 28])
shade, last = False, None
for row in ws.iter_rows(min_row=2):
    if row[0].value != last:
        shade, last = not shade, row[0].value
    if shade:
        for c in row:
            c.fill = ALT_FILL
ws.auto_filter.ref = ws.dimensions

# Sample data
ws = wb.create_sheet("Sample data")
ws.append(["Part", "What gets seeded"])
for k, v in SAMPLE_DATA.items():
    ws.append([k.capitalize(), v])
style_sheet(ws, [16, 110])

# Category counts
ws = wb.create_sheet("Coverage")
ws.append(["Category", "Scenarios", "Steps", "Multi-turn scenarios"])
cats = list(dict.fromkeys(s["category"] for s in SCENARIOS))
for c in cats:
    sc = [s for s in SCENARIOS if s["category"] == c]
    ws.append([c, len(sc), sum(len(s["steps"]) for s in sc), sum(1 for s in sc if len(s["steps"]) > 1)])
n = len(cats) + 1
ws.append(["Total", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})"])
style_sheet(ws, [22, 12, 10, 20])
for c in ws[n + 1]:
    c.font = Font(name=FONT, size=10, bold=True)

out = Path(__file__).resolve().parent / "test_plan.xlsx"
wb.save(out)
print(f"Wrote {out} · {len(SCENARIOS)} scenarios · {total_steps} steps · {multi} multi-turn")
