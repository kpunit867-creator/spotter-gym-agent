"""Writes evaluation results to tests/results/eval_<run>.xlsx (Summary, Steps, LLM calls, Errors)."""
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tests.metrics import STEP_METRICS

RESULTS_DIR = Path(__file__).resolve().parent / "results"
# Optional: path to a LibreOffice-based recalc script that fills in formula values after writing.
RECALC = Path(os.environ["XLSX_RECALC_SCRIPT"]) if os.getenv("XLSX_RECALC_SCRIPT") else None
FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E78")
GREEN = PatternFill("solid", fgColor="E2F0D9")
RED = PatternFill("solid", fgColor="FBE2E2")
WIDE = {"agent_reply": 60, "input_message": 40, "expected_behaviour": 36, "failed_checks": 50, "tool_args": 40,
        "tools_called": 30, "backend_error": 30, "scenario": 26}


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _body_style(ws, widths=None):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(cell.value, 14)


def write_results(rows, run_id):
    RESULTS_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    cols = [m[0] for m in STEP_METRICS]
    letter = {c: get_column_letter(i) for i, c in enumerate(cols, start=1)}
    n = len(rows) + 1  # last data row in Steps

    # ── Steps ──
    steps = wb.active
    steps.title = "Steps"
    _header(steps, cols)
    for r in rows:
        steps.append([r.get(c, "") for c in cols])
    _body_style(steps, WIDE)
    rng = f"{letter['status']}2:{letter['status']}{max(n, 2)}"
    steps.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN))
    steps.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED))
    steps.auto_filter.ref = steps.dimensions

    # ── LLM calls ──
    calls = wb.create_sheet("LLM calls")
    ccols = ["scenario_id", "step", "call_number", "model", "input_tokens", "output_tokens", "thinking_tokens",
             "cached_tokens", "latency_ms", "finish_reason"]
    _header(calls, ccols)
    for r in rows:
        for c in r.get("_calls", []):
            calls.append([r["scenario_id"], r["step"], c["call_number"], c["model_name"], c["input_tokens"],
                          c["output_tokens"], c["thinking_tokens"], c["cached_tokens"], c["latency_ms"], c["finish_reason"]])
    _body_style(calls)

    # ── Errors ──
    errs = wb.create_sheet("Errors")
    _header(errs, ["scenario_id", "step", "input_message", "backend_error", "tool_errors", "failed_checks"])
    for r in rows:
        if r["backend_error"] or r["tool_errors"] or r["status"] == "FAIL":
            errs.append([r["scenario_id"], r["step"], r["input_message"], r["backend_error"], r["tool_errors"], r["failed_checks"]])
    _body_style(errs, {"input_message": 40, "backend_error": 40, "tool_errors": 30, "failed_checks": 70})

    # ── Summary (formulas over Steps) ──
    s = wb.create_sheet("Summary", 0)
    col = lambda name: f"Steps!${letter[name]}$2:${letter[name]}${max(n, 2)}"  # noqa: E731
    s.append(["Spotter evaluation", f"Run {run_id}"])
    s.append([])
    s.append(["Overall", "Value", "Notes"])
    head_row = s.max_row
    overall = [
        ("Steps run", f"=COUNTA({col('scenario_id')})", ""),
        ("Steps passed", f'=COUNTIF({col("status")},"PASS")', ""),
        ("Pass rate", "=IFERROR({r:Steps passed}/{r:Steps run},0)", "passed ÷ run"),
        ("Avg total time (ms)", f"=AVERAGE({col('total_time_ms')})", "message in → reply saved"),
        ("Median total time (ms)", f"=MEDIAN({col('total_time_ms')})", ""),
        ("95th percentile time (ms)", f"=PERCENTILE({col('total_time_ms')},0.95)", "slowest 5% are above this"),
        ("Avg Gemini time (ms)", f"=AVERAGE({col('llm_time_ms')})", ""),
        ("Avg backend time (ms)", f"=AVERAGE({col('backend_time_ms')})", "our code + database (from this Mac)"),
        ("Avg DB queries per step", f"=AVERAGE({col('db_query_count')})", ""),
        ("Avg Gemini calls per step", f"=AVERAGE({col('llm_calls')})", "1.0 is ideal for plain logging"),
        ("Avg input tokens", f"=AVERAGE({col('input_tokens')})", ""),
        ("Avg output tokens", f"=AVERAGE({col('output_tokens')})", ""),
        ("Total cached tokens", f"=SUM({col('cached_tokens')})", "Gemini's implicit cache"),
        ("Total est. cost (USD)", f"=SUM({col('est_cost_usd')})", ""),
        ("Avg cost per message (USD)", f"=AVERAGE({col('est_cost_usd')})", ""),
        ("Projected monthly cost (USD)", "={r:Avg cost per message (USD)}*450", "at ~450 messages/month, one user"),
        ("Steps with backend errors", f'=COUNTIF({col("backend_error")},"?*")', ""),
        ("Steps with tool errors", f'=COUNTIF({col("tool_errors")},"?*")', "tool rejected the model's input"),
        ("Clarifying questions asked", f'=COUNTIF({col("outcome")},"needs_clarification")', ""),
        ("Rate-limit retries", f"=SUM({col('retries')})", ""),
    ]
    row_of = {label: head_row + 1 + i for i, (label, _, _) in enumerate(overall)}
    for label, formula, note in overall:
        for name, r in row_of.items():
            formula = formula.replace("{r:" + name + "}", f"B{r}")
        s.append([label, formula, note])
    s.append([])
    s.append(["By category", "Steps", "Passed", "Pass rate", "Avg time (ms)", "Avg Gemini calls"])
    cat_head = s.max_row
    for cat in dict.fromkeys(r["category"] for r in rows):
        i = s.max_row + 1
        s.append([cat, f'=COUNTIF({col("category")},A{i})',
                  f'=COUNTIFS({col("category")},A{i},{col("status")},"PASS")', f"=IFERROR(C{i}/B{i},0)",
                  f'=AVERAGEIF({col("category")},A{i},{col("total_time_ms")})',
                  f'=AVERAGEIF({col("category")},A{i},{col("llm_calls")})'])
    s.append([])
    s.append(["Checks", "Passed", "Failed", "Not applicable"])
    chk_head = s.max_row
    for chk in ("check_intent", "check_tools", "check_data", "check_reply", "check_style"):
        s.append([chk, f'=COUNTIF({col(chk)},"PASS")', f'=COUNTIF({col(chk)},"FAIL")', f'=COUNTIF({col(chk)},"n/a")'])

    for row in s.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row in (head_row, cat_head, chk_head) or c.row == 1)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    s["A1"].font = Font(name=FONT, size=14, bold=True)
    for r in (head_row, cat_head, chk_head):
        for c in s[r]:
            if c.value:
                c.fill = HEAD
                c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    formats = {"Pass rate": "0.0%", "Avg Gemini calls per step": "0.00", "Avg DB queries per step": "0.0",
               "Total est. cost (USD)": "$0.0000", "Avg cost per message (USD)": "$0.00000",
               "Projected monthly cost (USD)": "$0.00"}
    for label, r in row_of.items():
        s[f"B{r}"].number_format = formats.get(label, "#,##0")
    for r in range(cat_head + 1, chk_head - 1):
        s[f"D{r}"].number_format = "0.0%"
        s[f"E{r}"].number_format = "#,##0"
        s[f"F{r}"].number_format = "0.00"
    for letter_, w in zip("ABCDEF", (32, 16, 40, 12, 14, 16)):
        s.column_dimensions[letter_].width = w

    path = RESULTS_DIR / f"eval_{run_id}.xlsx"
    wb.save(path)
    if RECALC and RECALC.exists():
        subprocess.run([sys.executable, str(RECALC), str(path), "90"], capture_output=True)
    return path
